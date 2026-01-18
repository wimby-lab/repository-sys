import io
import shutil
import tempfile

from django.test import TestCase, Client, override_settings
from django.urls import reverse
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document as DocxDocument
from openpyxl import Workbook
from accounts.models import User, Role
from .models import Document, DocumentFolder
from .forms import DocumentFolderForm, DocumentSearchForm
from .permissions import can_access_document


class DocumentAccessTests(TestCase):
    """Test document access control"""
    
    def setUp(self):
        self.client = Client()
        
        # Create roles
        self.admin_role = Role.objects.create(name=Role.ADVISER)
        self.manager_role = Role.objects.create(name=Role.PRESIDENT)
        self.user_role = Role.objects.create(name=Role.AUDITOR)
        
        # Create users
        self.admin = User.objects.create_user(
            username='admin',
            password='pass',
            role=self.admin_role
        )
        self.manager = User.objects.create_user(
            username='manager',
            password='pass',
            role=self.manager_role
        )
        self.user1 = User.objects.create_user(
            username='user1',
            password='pass',
            role=self.user_role
        )
        self.user2 = User.objects.create_user(
            username='user2',
            password='pass',
            role=self.user_role
        )
        
        # Create test documents
        self.public_doc = Document.objects.create(
            title='Public Document',
            owner=self.user1,
            classification='PUBLIC',
            file='test.txt',
            file_type='text/plain',
            file_size=100
        )
        
        self.internal_doc = Document.objects.create(
            title='Internal Document',
            owner=self.user1,
            classification='INTERNAL',
            file='test.txt',
            file_type='text/plain',
            file_size=100
        )
        
        self.restricted_doc = Document.objects.create(
            title='Restricted Document',
            owner=self.user1,
            classification='RESTRICTED',
            file='test.txt',
            file_type='text/plain',
            file_size=100
        )
        
    def test_owner_can_access_own_documents(self):
        """Test document owner can access their own documents"""
        self.assertTrue(can_access_document(self.user1, self.public_doc))
        self.assertTrue(can_access_document(self.user1, self.internal_doc))
        self.assertTrue(can_access_document(self.user1, self.restricted_doc))
        
    def test_adviser_can_access_all_documents(self):
        """Test adviser can access all documents"""
        self.assertTrue(can_access_document(self.admin, self.public_doc))
        self.assertTrue(can_access_document(self.admin, self.internal_doc))
        self.assertTrue(can_access_document(self.admin, self.restricted_doc))
        
    def test_president_cannot_access_restricted(self):
        """Test president cannot access restricted documents"""
        self.assertTrue(can_access_document(self.manager, self.public_doc))
        self.assertTrue(can_access_document(self.manager, self.internal_doc))
        self.assertFalse(can_access_document(self.manager, self.restricted_doc))
        
    def test_officer_can_access_public_only(self):
        """Test officer can only access public documents by default"""
        self.assertTrue(can_access_document(self.user2, self.public_doc))
        self.assertFalse(can_access_document(self.user2, self.internal_doc))
        self.assertFalse(can_access_document(self.user2, self.restricted_doc))
        
    def test_shared_document_access(self):
        """Test user can access documents shared with them"""
        self.internal_doc.shared_with.add(self.user2)
        self.assertTrue(can_access_document(self.user2, self.internal_doc))
        
    def test_document_upload_requires_login(self):
        """Test document upload requires authentication"""
        response = self.client.get(reverse('documents:document_upload'))
        self.assertEqual(response.status_code, 302)  # Redirect to login
        
    def test_authenticated_user_can_upload(self):
        """Test authenticated user can access upload page"""
        self.client.login(username='user1', password='pass')
        response = self.client.get(reverse('documents:document_upload'))
        self.assertEqual(response.status_code, 200)

    def test_upload_sets_file_metadata(self):
        """Test upload stores file metadata from the uploaded file"""
        self.client.login(username='user1', password='pass')
        upload = SimpleUploadedFile(
            'hello.txt',
            b'hello',
            content_type='text/plain'
        )

        response = self.client.post(
            reverse('documents:document_upload'),
            {
                'title': 'Hello',
                'description': 'Test upload',
                'file': upload,
                'classification': 'PUBLIC',
                'section': 'GENERAL',
                'category': 'General',
                'tags': 'test',
            }
        )

        self.assertEqual(response.status_code, 302)
        document = Document.objects.get(title='Hello')
        self.assertEqual(document.file_size, 5)
        self.assertEqual(document.file_type, 'text/plain')

    def test_upload_allows_google_docs_link(self):
        """Test upload works with a Google Docs link"""
        self.client.login(username='user1', password='pass')
        google_link = 'https://docs.google.com/document/d/abc123/edit'

        response = self.client.post(
            reverse('documents:document_upload'),
            {
                'title': 'Linked Doc',
                'description': 'Google Docs link',
                'google_docs_url': google_link,
                'classification': 'PUBLIC',
                'section': 'GENERAL',
                'category': 'General',
                'tags': 'google',
            }
        )

        self.assertEqual(response.status_code, 302)
        document = Document.objects.get(title='Linked Doc')
        self.assertEqual(document.google_docs_url, google_link)
        self.assertEqual(document.file_type, 'Google Docs')
        self.assertEqual(document.file_size, 0)

    def test_upload_rejects_invalid_google_docs_link(self):
        """Test upload rejects non-Google Docs URLs"""
        self.client.login(username='user1', password='pass')
        response = self.client.post(
            reverse('documents:document_upload'),
            {
                'title': 'Invalid Doc',
                'description': 'Bad link',
                'google_docs_url': 'https://example.com/document/123',
                'classification': 'PUBLIC',
                'section': 'GENERAL',
                'category': 'General',
                'tags': 'google',
            }
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Google Docs URL must be a valid docs.google.com document link.')
        
    def test_archived_documents_not_in_list(self):
        """Test archived documents are not shown in document list"""
        # Archive a document
        self.public_doc.is_archived = True
        self.public_doc.archived_by = self.admin
        self.public_doc.save()
        
        # Login and get document list
        self.client.login(username='user2', password='pass')
        response = self.client.get(reverse('documents:document_list'))
        
        # Archived document should not be in the list
        self.assertNotContains(response, self.public_doc.title)
        
    def test_archived_document_detail_blocked(self):
        """Test archived documents cannot be accessed via detail view"""
        # Archive a document
        self.public_doc.is_archived = True
        self.public_doc.archived_by = self.admin
        self.public_doc.save()
        
        # Try to access archived document
        self.client.login(username='user2', password='pass')
        response = self.client.get(reverse('documents:document_detail', args=[self.public_doc.pk]))
        
        # Should redirect with error message
        self.assertEqual(response.status_code, 302)
        
    def test_archived_document_download_blocked(self):
        """Test archived documents cannot be downloaded"""
        # Archive a document
        self.public_doc.is_archived = True
        self.public_doc.archived_by = self.admin
        self.public_doc.save()
        
        # Try to download archived document
        self.client.login(username='user2', password='pass')
        response = self.client.get(reverse('documents:document_download', args=[self.public_doc.pk]))
        
        # Should redirect with error message
        self.assertEqual(response.status_code, 302)

    def test_document_list_filters_by_section(self):
        """Test document list can be filtered by section"""
        folder = DocumentFolder.objects.get(key='GENERAL')
        section_value = folder.key
        section_doc = Document.objects.create(
            title='Section Document',
            owner=self.user1,
            classification='PUBLIC',
            section=section_value,
            file='test.txt',
            file_type='text/plain',
            file_size=100
        )

        self.client.login(username='user1', password='pass')
        response = self.client.get(reverse('documents:document_list'), {'section': section_value})

        self.assertContains(response, section_doc.title)
        self.assertNotContains(response, self.public_doc.title)


class DocumentFolderTests(TestCase):
    """Test document folder management"""

    def test_folder_form_generates_key(self):
        form = DocumentFolderForm(data={'name': 'Project Docs'})
        self.assertTrue(form.is_valid())
        folder = form.save()
        self.assertEqual(folder.key, 'PROJECT_DOCS')

    def test_search_form_includes_folder_choices(self):
        folder = DocumentFolder.objects.create(key='FINANCE', name='Finance')
        form = DocumentSearchForm()
        self.assertIn((folder.key, folder.name), form.fields['section'].choices)


class DocumentPreviewTests(TestCase):
    """Test document preview rendering"""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.temp_media_root = tempfile.mkdtemp()
        cls.override_media = override_settings(MEDIA_ROOT=cls.temp_media_root)
        cls.override_media.enable()

    @classmethod
    def tearDownClass(cls):
        cls.override_media.disable()
        shutil.rmtree(cls.temp_media_root, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.user_role = Role.objects.create(name=Role.AUDITOR)
        self.user = User.objects.create_user(
            username='preview_user',
            password='pass',
            role=self.user_role
        )
        self.client = Client()
        self.client.login(username='preview_user', password='pass')

    def test_text_document_preview(self):
        upload = SimpleUploadedFile(
            'preview.txt',
            b'Preview line one\nPreview line two',
            content_type='text/plain'
        )
        document = Document.objects.create(
            title='Text Preview',
            owner=self.user,
            classification='PUBLIC',
            section='GENERAL',
            file=upload,
            file_type=upload.content_type,
            file_size=upload.size
        )

        response = self.client.get(reverse('documents:document_detail', args=[document.pk]))

        self.assertContains(response, 'Preview line one')
        self.assertContains(response, 'Preview line two')

    def test_docx_document_preview(self):
        buffer = io.BytesIO()
        doc = DocxDocument()
        doc.add_paragraph('Docx preview content')
        doc.save(buffer)
        buffer.seek(0)
        upload = SimpleUploadedFile(
            'preview.docx',
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        document = Document.objects.create(
            title='Docx Preview',
            owner=self.user,
            classification='PUBLIC',
            section='GENERAL',
            file=upload,
            file_type=upload.content_type,
            file_size=upload.size
        )

        response = self.client.get(reverse('documents:document_detail', args=[document.pk]))

        self.assertContains(response, 'Docx preview content')

    def test_spreadsheet_document_preview(self):
        buffer = io.BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet['A1'] = 'Header'
        sheet['A2'] = 'Value'
        workbook.save(buffer)
        buffer.seek(0)
        upload = SimpleUploadedFile(
            'preview.xlsx',
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        document = Document.objects.create(
            title='Spreadsheet Preview',
            owner=self.user,
            classification='PUBLIC',
            section='GENERAL',
            file=upload,
            file_type=upload.content_type,
            file_size=upload.size
        )

        response = self.client.get(reverse('documents:document_detail', args=[document.pk]))

        self.assertContains(response, 'Header')
        self.assertContains(response, 'Value')


class DocumentModelTests(TestCase):
    """Test document model"""
    
    def setUp(self):
        self.user_role = Role.objects.create(name=Role.SECRETARY)
        self.user = User.objects.create_user(
            username='testuser',
            password='pass',
            role=self.user_role
        )
        
    def test_document_creation(self):
        """Test document can be created"""
        doc = Document.objects.create(
            title='Test Document',
            owner=self.user,
            classification='PUBLIC',
            file='test.txt',
            file_type='text/plain',
            file_size=100
        )
        self.assertEqual(doc.title, 'Test Document')
        self.assertEqual(doc.owner, self.user)
        self.assertEqual(doc.section, Document.SECTION_CHOICES[0][0])
        self.assertFalse(doc.is_archived)
        
    def test_document_tags_list(self):
        """Test document tags parsing"""
        doc = Document.objects.create(
            title='Test Document',
            owner=self.user,
            classification='PUBLIC',
            tags='tag1, tag2, tag3',
            file='test.txt',
            file_type='text/plain',
            file_size=100
        )
        tags = doc.get_tags_list()
        self.assertEqual(len(tags), 3)
        self.assertIn('tag1', tags)
        self.assertIn('tag2', tags)
        self.assertIn('tag3', tags)
        
    def test_document_archiving(self):
        """Test document can be archived"""
        from django.utils import timezone
        
        doc = Document.objects.create(
            title='Test Document',
            owner=self.user,
            classification='PUBLIC',
            file='test.txt',
            file_type='text/plain',
            file_size=100
        )
        self.assertFalse(doc.is_archived)
        
        # Archive the document
        doc.is_archived = True
        doc.archived_by = self.user
        doc.archived_at = timezone.now()
        doc.save()
        
        self.assertTrue(doc.is_archived)
        self.assertEqual(doc.archived_by, self.user)
        self.assertIsNotNone(doc.archived_at)
